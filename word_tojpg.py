import os
import tempfile
from pathlib import Path

import fitz  # PyMuPDF
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


# =========================
# 你只需要改这里
# =========================

# 存放 PDF 的文件夹路径
PDF_FOLDER = Path(r"D:\Desktop\财务工作\智能报销\2025 7A  5265")

# 输出 Word 文件路径
OUTPUT_WORD = POUTPUT_WORD = Path(r"D:\Desktop\财务工作\智能报销\2025 7A  5265.docx")
# 图片清晰度，2.0 一般够用；越大越清晰，但 Word 文件也越大
ZOOM = 2.0

# 插入 Word 的图片宽度，单位是英寸
IMAGE_WIDTH_INCHES = 6.5


def get_pdf_files(folder: Path):
    """
    获取指定文件夹下的所有 PDF 文件。
    这里只读取当前文件夹，不读取子文件夹。
    """
    pdf_files = []

    for file in folder.iterdir():
        if file.is_file() and file.suffix.lower() == ".pdf":
            pdf_files.append(file)

    return sorted(pdf_files, key=lambda x: x.name)


def pdf_pages_to_images(pdf_path: Path, temp_folder: Path):
    """
    将 PDF 的每一页导出为图片。
    返回图片路径列表。
    """
    image_paths = []

    pdf_doc = fitz.open(str(pdf_path))

    for page_index in range(len(pdf_doc)):
        page = pdf_doc.load_page(page_index)

        matrix = fitz.Matrix(ZOOM, ZOOM)
        pix = page.get_pixmap(matrix=matrix, alpha=False)

        image_path = temp_folder / f"{pdf_path.stem}_第{page_index + 1:03d}页.png"
        pix.save(str(image_path))

        image_paths.append(image_path)

    pdf_doc.close()

    return image_paths


def add_pdf_to_word(document: Document, pdf_path: Path, image_paths):
    """
    将一个 PDF 对应的所有页面图片插入 Word。
    """
    document.add_heading(pdf_path.name, level=1)

    for index, image_path in enumerate(image_paths, start=1):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = paragraph.add_run()
        run.add_picture(str(image_path), width=Inches(IMAGE_WIDTH_INCHES))

        caption = document.add_paragraph(f"{pdf_path.name} - 第 {index} 页")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # 每页图片后分页，最后一页也分页没关系
        document.add_page_break()


def convert_all_pdfs_to_word(pdf_folder: Path, output_word: Path):
    """
    主函数：读取文件夹内所有 PDF，将内容转为图片后写入 Word。
    """
    if not pdf_folder.exists():
        print(f"错误：文件夹不存在：{pdf_folder}")
        return

    if not pdf_folder.is_dir():
        print(f"错误：这不是文件夹：{pdf_folder}")
        return

    pdf_files = get_pdf_files(pdf_folder)

    if not pdf_files:
        print(f"没有找到 PDF 文件：{pdf_folder}")
        return

    output_word.parent.mkdir(parents=True, exist_ok=True)

    document = Document()

    document.add_heading("PDF 内容汇总", level=0)
    document.add_paragraph(f"来源文件夹：{pdf_folder}")
    document.add_paragraph(f"PDF 数量：{len(pdf_files)}")
    document.add_page_break()

    with tempfile.TemporaryDirectory() as temp_dir:
        temp_folder = Path(temp_dir)

        for pdf_file in pdf_files:
            print(f"正在处理：{pdf_file.name}")

            try:
                image_paths = pdf_pages_to_images(pdf_file, temp_folder)
                add_pdf_to_word(document, pdf_file, image_paths)

            except Exception as e:
                print(f"处理失败：{pdf_file}")
                print(f"错误信息：{e}")

                document.add_heading(f"{pdf_file.name} 处理失败", level=1)
                document.add_paragraph(f"错误信息：{e}")
                document.add_page_break()

    document.save(str(output_word))
    print(f"\n完成！Word 已保存到：{output_word}")


if __name__ == "__main__":
    convert_all_pdfs_to_word(PDF_FOLDER, OUTPUT_WORD)